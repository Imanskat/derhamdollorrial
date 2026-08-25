#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ساخت فایل اکسل «حساب‌کتاب واردات».

همان مدل ساده داشبورد وب، به‌صورت فرمول‌های زنده:
    قیمت خرید (به هر ارزی)  +  هزینه آوردن تا ایران  =  بهای تمام‌شده

هیچ عددی در ستون‌های محاسباتی ثابت نوشته نشده؛ همه از سلول‌های ورودی می‌آیند.

اجرا:  python3 build_workbook.py
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT   = "حساب-کتاب-واردات.xlsx"
FONT  = "Tahoma"
NROWS = 25

BLUE  = Font(name=FONT, size=10, color="0000FF")
BLACK = Font(name=FONT, size=10)
GREEN = Font(name=FONT, size=10, color="008000")
BOLD  = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True)
SECT  = Font(name=FONT, size=11, bold=True, color="1F3864")
SMALL = Font(name=FONT, size=8.5, color="595959")

YELLOW   = PatternFill("solid", fgColor="FFF2A8")
HEADFILL = PatternFill("solid", fgColor="1F3864")
HEADFONT = Font(name=FONT, size=9, bold=True, color="FFFFFF")
GREYFILL = PatternFill("solid", fgColor="F2F2F2")
TOTFILL  = PatternFill("solid", fgColor="E2EFDA")

thin = Side(style="thin", color="BFBFBF")
BOX  = Border(left=thin, right=thin, top=thin, bottom=thin)

TOM  = '#,##0;(#,##0);-'
CUR2 = '#,##0.00;(#,##0.00);-'
PCT  = '0.0%;(0.0%);-'
NUM  = '#,##0;(#,##0);-'
NUM1 = '#,##0.0;(#,##0.0);-'
RATE = '#,##0.0000'

# ارزها: (نام فارسی، چند واحد به ازای یک دلار)
CURRENCIES = [("دلار",1.0), ("درهم",3.6725), ("یوان",7.10),
              ("یورو",0.92), ("لیر",34.0), ("پوند",0.79), ("روبل",90.0)]
COST_CUR   = [c[0] for c in CURRENCIES] + ["تومان"]

wb = Workbook()

def put(ws, cell, value, font=BLACK, fmt=None, fill=None, border=False, align=None):
    c = ws[cell]; c.value = value; c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if border: c.border = BOX
    if align: c.alignment = Alignment(horizontal=align, vertical="center")
    return c

# ======================================================================
#  شیت ۱ — ارز و هزینه‌ها
# ======================================================================
p = wb.active
p.title = "ارز و هزینه‌ها"
p.sheet_view.rightToLeft = True
p.column_dimensions["A"].width = 34
p.column_dimensions["B"].width = 18
p.column_dimensions["C"].width = 18
p.column_dimensions["D"].width = 50

put(p, "A1", "حساب‌کتاب واردات — ارز و هزینه‌ها", TITLE)
put(p, "A2", "فقط خانه‌های زرد را پر کنید. بقیه خودکار حساب می‌شوند.", SMALL)

put(p, "A4", "۱ — ارز", SECT)
put(p, "A5", "قیمت دلار (تومان)", BOLD)
put(p, "B5", 115_000, BLUE, TOM, YELLOW, True)
put(p, "D5", "نرخ بازار آزاد امروز. تنها نرخی که باید دستی وارد کنید.", SMALL)

put(p, "A7", "قیمت بقیه ارزها از روی دلار حساب می‌شود:", SMALL)
for j, t in enumerate(["ارز", "واحد به ازای ۱ دلار", "قیمت هر واحد (تومان)"]):
    c = p.cell(row=8, column=1+j, value=t)
    c.font = HEADFONT; c.fill = HEADFILL; c.border = BOX
    c.alignment = Alignment(horizontal="center", vertical="center")

CUR_FIRST = 9
for i, (fa, per) in enumerate(CURRENCIES):
    r = CUR_FIRST + i
    put(p, f"A{r}", fa, BLACK, None, None, True)
    put(p, f"B{r}", per, BLUE, RATE, YELLOW, True)
    put(p, f"C{r}", f"=IF($B{r}=0,0,$B$5/$B{r})", BLACK, TOM, GREYFILL, True)
CUR_LAST = CUR_FIRST + len(CURRENCIES) - 1

TOMAN_ROW = CUR_LAST + 1
put(p, f"A{TOMAN_ROW}", "تومان", BLACK, None, None, True)
put(p, f"B{TOMAN_ROW}", 1, BLACK, RATE, GREYFILL, True)
put(p, f"C{TOMAN_ROW}", 1, BLACK, TOM, GREYFILL, True)
put(p, f"D{TOMAN_ROW}", "برای هزینه‌هایی که به تومان پرداخت شده‌اند.", SMALL)
LOOK_FIRST, LOOK_LAST = CUR_FIRST, TOMAN_ROW

r = TOMAN_ROW + 2
put(p, f"A{r}", "۲ — هزینه آوردن جنس تا ایران", SECT); r += 1
put(p, f"A{r}", "این مبالغ مال کل محموله است، نه هر کالا.", SMALL); r += 1

for j, t in enumerate(["هزینه", "مبلغ", "ارز"]):
    c = p.cell(row=r, column=1+j, value=t)
    c.font = HEADFONT; c.fill = HEADFILL; c.border = BOX
    c.alignment = Alignment(horizontal="center", vertical="center")
r += 1
COST_ROWS = {}
for name, amt, cur, note in [
    ("کرایه حمل تا ایران", 1800,          "دلار",  "دریایی، زمینی یا هوایی — مبلغ کل"),
    ("ترخیص و گمرک",      1_400_000_000, "تومان", "هرچه بابت گمرک و ترخیص‌کار دادید، یک‌جا"),
    ("سایر هزینه‌ها",      95_000_000,    "تومان", "انبار، حمل داخلی، بسته‌بندی، مجوز"),
]:
    put(p, f"A{r}", name, BLACK, None, None, True)
    put(p, f"B{r}", amt, BLUE, TOM, YELLOW, True)
    put(p, f"C{r}", cur, BLUE, None, YELLOW, True)
    put(p, f"D{r}", note, SMALL)
    COST_ROWS[name] = r
    r += 1

FREIGHT_R = COST_ROWS["کرایه حمل تا ایران"]
CUSTOMS_R = COST_ROWS["ترخیص و گمرک"]
OTHER_R   = COST_ROWS["سایر هزینه‌ها"]

r += 1
SPLIT_R = r
put(p, f"A{r}", "تقسیم هزینه بین کالاها بر اساس", BLACK, None, None, True)
put(p, f"B{r}", "قیمت کالا", BLUE, None, YELLOW, True)
put(p, f"D{r}", "«قیمت کالا» یا «وزن کالا» — از فهرست انتخاب کنید.", SMALL); r += 1

TARGET_R = r
put(p, f"A{r}", "سود دلخواه (٪)", BLACK, None, None, True)
put(p, f"B{r}", 0.35, BLUE, PCT, YELLOW, True)
put(p, f"D{r}", "مبنای ستون «قیمت پیشنهادی» در شیت کالاها.", SMALL); r += 2

put(p, f"A{r}", "راهنمای رنگ‌ها", SECT); r += 1
put(p, f"A{r}", "زرد = شما پر می‌کنید.", BLUE); r += 1
put(p, f"A{r}", "خاکستری = فرمول؛ دست نزنید.", BLACK); r += 2
put(p, f"A{r}", "قیمت دلار را از صرافی بگیرید. نرخ رسمی یا نیما را وارد نکنید — "
                "بهای تمام‌شده چند برابر کمتر از واقعیت درمی‌آید.", SMALL)

dv_cost = DataValidation(type="list", formula1='"' + ",".join(COST_CUR) + '"', allow_blank=False)
p.add_data_validation(dv_cost)
for rr in COST_ROWS.values():
    dv_cost.add(p[f"C{rr}"])
dv_split = DataValidation(type="list", formula1='"قیمت کالا,وزن کالا"', allow_blank=False)
p.add_data_validation(dv_split); dv_split.add(p[f"B{SPLIT_R}"])

# ======================================================================
#  شیت ۲ — کالاها
# ======================================================================
s = wb.create_sheet("کالاها")
s.sheet_view.rightToLeft = True

P = "'ارز و هزینه‌ها'!"
USD      = f"{P}$B$5"
LOOK_A   = f"{P}$A${LOOK_FIRST}:$A${LOOK_LAST}"
LOOK_C   = f"{P}$C${LOOK_FIRST}:$C${LOOK_LAST}"
SPLIT    = f"{P}$B${SPLIT_R}"
TARGET   = f"{P}$B${TARGET_R}"

def cost_toman(row):
    """مبلغ یک هزینه، تبدیل‌شده به تومان."""
    return (f"{P}$B${row}*INDEX({LOOK_C},MATCH({P}$C${row},{LOOK_A},0))")

FREIGHT_T = cost_toman(FREIGHT_R)
CUSTOMS_T = cost_toman(CUSTOMS_R)
OTHER_T   = cost_toman(OTHER_R)

COLS = [
    ("نام کالا",                  24, None, "in"),
    ("ارز خرید",                  11, None, "in"),
    ("قیمت خرید\nهر عدد",         13, CUR2, "in"),
    ("تعداد",                     10, NUM,  "in"),
    ("وزن کل\n(کیلوگرم)",         12, NUM1, "in"),
    ("قیمت فروش\nهر عدد (تومان)", 15, TOM,  "in"),
    ("قیمت ارز\n(تومان)",         13, TOM,  "fx"),
    ("ارزش خرید\n(تومان)",        16, TOM,  "fx"),
    ("سهم از هزینه‌ها",           13, PCT,  "fx"),
    ("کرایه حمل\n(تومان)",        14, TOM,  "fx"),
    ("ترخیص و گمرک\n(تومان)",     15, TOM,  "fx"),
    ("سایر\n(تومان)",             13, TOM,  "fx"),
    ("بهای تمام‌شده کل\n(تومان)", 17, TOM,  "fx"),
    ("بهای تمام‌شده\nهر عدد",     15, TOM,  "fx"),
    ("سود هر عدد\n(تومان)",       14, TOM,  "fx"),
    ("درصد سود",                  11, PCT,  "fx"),
    ("سود کل\n(تومان)",           16, TOM,  "fx"),
    ("کمترین قیمت\nبدون ضرر",     15, TOM,  "fx"),
    ("قیمت پیشنهادی",             15, TOM,  "fx"),
]

put(s, "A1", "کالاها", TITLE)
put(s, "A2", "ستون‌های زرد (A تا F) را پر کنید؛ بقیه خودکار حساب می‌شوند.", SMALL)

HEAD = 4
FIRST = HEAD + 1
LAST  = FIRST + NROWS - 1
TOT   = LAST + 1

for i, (t, w, fmt, kind) in enumerate(COLS, start=1):
    L = get_column_letter(i)
    s.column_dimensions[L].width = w
    c = s.cell(row=HEAD, column=i, value=t)
    c.font = HEADFONT; c.fill = HEADFILL; c.border = BOX
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
s.row_dimensions[HEAD].height = 40
s.freeze_panes = "B5"

SAMPLE = [
    ("کرم پودر",      "درهم", 32,  600, 42, 2_610_000),
    ("رژ لب",         "درهم", 14, 1200, 30, 1_140_000),
    ("پالت سایه",     "دلار", 13,  300, 39, 3_890_000),
    ("سرم ویتامین ث", "دلار", 15,  400, 32, 4_490_000),
    ("ریمل",          "یوان", 35,  900, 27, 1_470_000),
]

W_SUM = f"SUM($E${FIRST}:$E${LAST})"
V_SUM = f"SUM($H${FIRST}:$H${LAST})"

for r in range(FIRST, LAST + 1):
    g = 'IF($A{}="","",{{}})'.format(r)
    n = r - FIRST
    vals = SAMPLE[n] if n < len(SAMPLE) else ("", None, None, None, None, None)

    for i in range(1, 7):
        c = s.cell(row=r, column=i, value=vals[i-1])
        c.font = BLUE; c.fill = YELLOW; c.border = BOX
        if COLS[i-1][2]: c.number_format = COLS[i-1][2]

    F = {
        7:  f"INDEX({LOOK_C},MATCH($B{r},{LOOK_A},0))",              # قیمت ارز
        8:  f"$C{r}*$D{r}*$G{r}",                                    # ارزش خرید
        9:  (f'IF({SPLIT}="وزن کالا",'
             f'IF({W_SUM}=0,0,$E{r}/{W_SUM}),'
             f'IF({V_SUM}=0,0,$H{r}/{V_SUM}))'),                     # سهم
        10: f"{FREIGHT_T}*$I{r}",                                    # کرایه
        11: f"{CUSTOMS_T}*$I{r}",                                    # ترخیص
        12: f"{OTHER_T}*$I{r}",                                      # سایر
        13: f"$H{r}+$J{r}+$K{r}+$L{r}",                              # بهای تمام‌شده کل
        14: f"IF($D{r}=0,0,$M{r}/$D{r})",                            # هر عدد
        15: f"$F{r}-$N{r}",                                          # سود هر عدد
        16: f"IF($F{r}=0,0,$O{r}/$F{r})",                            # درصد سود
        17: f"$O{r}*$D{r}",                                          # سود کل
        18: f"$N{r}",                                                # کمترین قیمت بدون ضرر
        19: f"IF(1-{TARGET}<=0,0,$N{r}/(1-{TARGET}))",               # قیمت پیشنهادی
    }
    for i, body in F.items():
        c = s.cell(row=r, column=i, value="=" + g.format(body))
        c.font = BLACK; c.fill = GREYFILL; c.border = BOX
        if COLS[i-1][2]: c.number_format = COLS[i-1][2]

# ردیف جمع
put(s, f"A{TOT}", "جمع", BOLD, None, TOTFILL, True)
SUMMABLE = {4, 5, 8, 9, 10, 11, 12, 13, 17}
for i in range(2, len(COLS) + 1):
    L = get_column_letter(i)
    c = s.cell(row=TOT, column=i)
    c.font = BOLD; c.fill = TOTFILL; c.border = BOX
    if COLS[i-1][2]: c.number_format = COLS[i-1][2]
    if i in SUMMABLE:
        c.value = f"=SUM({L}{FIRST}:{L}{LAST})"
    elif i == 16:
        rev = f"SUMPRODUCT($F${FIRST}:$F${LAST},$D${FIRST}:$D${LAST})"
        c.value = f"=IF({rev}=0,0,$Q${TOT}/{rev})"
    else:
        c.value = None

dv_row = DataValidation(type="list", formula1='"' + ",".join(c[0] for c in CURRENCIES) + '"',
                        allow_blank=True)
s.add_data_validation(dv_row)
for r in range(FIRST, LAST + 1):
    dv_row.add(s[f"B{r}"])

put(s, f"A{TOT+2}",
    "«کمترین قیمت بدون ضرر» همان بهای تمام‌شده هر عدد است — پایین‌تر از آن یعنی ضرر.", SMALL)
put(s, f"A{TOT+3}",
    "«قیمت پیشنهادی» قیمتی است که سود دلخواه شیت اول را می‌دهد.", SMALL)

# ======================================================================
#  شیت ۳ — راهنما
# ======================================================================
h = wb.create_sheet("راهنما")
h.sheet_view.rightToLeft = True
h.column_dimensions["A"].width = 28
h.column_dimensions["B"].width = 96
put(h, "A1", "راهنما", TITLE)

GUIDE = [
    ("۱ — قیمت دلار",
     "در شیت «ارز و هزینه‌ها» خانه B5 را با نرخ بازار آزاد امروز پر کنید. "
     "قیمت درهم، یوان، یورو و بقیه ارزها خودشان از روی همین حساب می‌شوند، چون نسبتشان با دلار تقریباً ثابت است."),
    ("۲ — هزینه آوردن",
     "کرایه حمل، ترخیص و گمرک، و سایر هزینه‌ها را وارد کنید. هرکدام را می‌توانید به ارز خودش بنویسید. "
     "این مبالغ مال کل محموله است و ابزار خودش بین کالاها تقسیم می‌کند."),
    ("۳ — کالاها",
     "در شیت «کالاها» ستون‌های A تا F را پر کنید: نام، ارز خرید، قیمت هر عدد، تعداد، وزن کل، قیمت فروش. "
     "تا ۲۵ کالا آماده است و ردیف‌های خالی نادیده گرفته می‌شوند."),
    ("", ""),
    ("بهای تمام‌شده یعنی چه",
     "قیمت خرید + سهم این کالا از کرایه حمل + سهم از ترخیص و گمرک + سهم از سایر هزینه‌ها. "
     "یعنی هر عدد از این کالا تا وقتی در انبار شما در ایران باشد، چند تمام شده."),
    ("تقسیم هزینه‌ها",
     "هزینه‌های کل محموله باید بین کالاها پخش شود. با «قیمت کالا» جنس گران‌تر سهم بیشتری می‌گیرد، "
     "با «وزن کالا» جنس سنگین‌تر. اگر بیشتر هزینه‌تان کرایه است وزن را انتخاب کنید، اگر گمرک است قیمت را."),
    ("قیمت پیشنهادی",
     "سود دلخواه را در شیت اول بگذارید؛ این ستون می‌گوید برای رسیدن به آن سود، هر عدد را چند بفروشید."),
    ("", ""),
    ("هشدار",
     "همه اعداد پیش‌فرض نمونه‌اند. قیمت دلار را از صرافی و هزینه ترخیص را از ترخیص‌کار خودتان بگیرید. "
     "نرخ رسمی یا نیما را وارد نکنید."),
]
rr = 3
for a, b in GUIDE:
    if a:
        put(h, f"A{rr}", a, BOLD)
        c = put(h, f"B{rr}", b, Font(name=FONT, size=10))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        h.row_dimensions[rr].height = 32
    rr += 1

wb.save(OUT)
print("نوشته شد:", OUT)
